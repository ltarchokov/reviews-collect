"""
Книга Excel с результатом — то, что уходит коллегам.

Четыре листа, всегда в одном порядке: «Сводка», «Отзывы», «Негатив»,
«Карточки». Порядок и вид одинаковы для всех площадок и всех запросов:
человек, получивший вторую книгу, ориентируется в ней без объяснений.

CSV рядом остаются: они для склейки датасетов, у них канонические
английские имена колонок. В Excel шапка русская — книгу читают,
а не мержат.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .doc import Doc, plain
from .schema import FIELDS, PRODUCT_FIELDS

LABELS = {
    "platform": "Площадка",
    "query": "Запрос",
    "brand": "Бренд",
    "product_name": "Товар",
    "product_id": "ID карточки",
    "product_url": "Ссылка на товар",
    "variant": "Вариант",
    "review_id": "ID отзыва",
    "author_hash": "Автор (хеш)",
    "rating": "Оценка",
    "date": "Дата",
    "title": "Заголовок",
    "text": "Текст",
    "pros": "Достоинства",
    "cons": "Недостатки",
    "has_photo": "Фото",
    "has_video": "Видео",
    "media_urls": "Медиа",
    "votes_up": "Полезно",
    "votes_down": "Бесполезно",
    "views": "Просмотры",
    "seller_answered": "Ответ продавца",
    "verified_purchase": "Покупка подтверждена",
    "tags": "Теги площадки",
    "about": "О чём",
    "source_url": "Ссылка на отзыв",
    "collected_at": "Собрано",
    "extra": "Дополнительно (JSON)",
    "supplier": "Продавец",
    "reviews_declared": "Отзывов заявлено",
    "reviews_collected": "Отзывов собрано",
    "price_rub": "Цена, ₽",
}

# Ширина колонок в символах. Текстовые поля широкие с переносом строк,
# служебные узкие. Не автоширина: по длинному отзыву она даёт колонку
# на весь экран, читать такое нельзя.
WIDTHS = {
    "text": 70, "pros": 36, "cons": 36, "title": 30, "product_name": 40,
    "product_url": 18, "source_url": 18, "brand": 14, "variant": 14,
    "tags": 24, "date": 11, "rating": 8, "platform": 10, "query": 14,
    "extra": 22, "supplier": 22, "media_urls": 14,
}
WRAP = {"text", "pros", "cons", "title", "product_name"}
URLS = {"product_url", "source_url"}

HEAD_FONT = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="E8E8E8")
H1 = Font(bold=True, size=14)
H2 = Font(bold=True, size=12)
H3 = Font(bold=True, size=11)
QUOTE = Font(italic=True, color="555555")
LINK = Font(color="0563C1", underline="single")


def _cell_value(field: str, value):
    if isinstance(value, bool):
        return "да" if value else ""
    if field == "rating" and str(value).isdigit():
        return int(value)
    if field in ("votes_up", "votes_down", "views", "reviews_declared", "reviews_collected"):
        try:
            return int(value)
        except (TypeError, ValueError):
            return value
    if field == "price_rub":
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    return value if value not in (None, "") else None


def _sheet_table(ws, rows: list[dict], fields: list[str]) -> None:
    """Лист-таблица: русская шапка, фильтр, закреплённая шапка, ширины."""
    ws.append([LABELS.get(f, f) for f in fields])
    for c in ws[1]:
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
        c.alignment = Alignment(vertical="top", wrap_text=True)
    for r in rows:
        ws.append([_cell_value(f, r.get(f)) for f in fields])
    for i, f in enumerate(fields, 1):
        col = get_column_letter(i)
        ws.column_dimensions[col].width = WIDTHS.get(f, 12)
        if f in WRAP:
            for c in ws[col][1:]:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        if f in URLS:
            for c in ws[col][1:]:
                if c.value:
                    c.hyperlink = str(c.value)
                    c.value = "открыть"
                    c.font = LINK
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.alignment.wrap_text is not True:
                c.alignment = Alignment(vertical="top")
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(fields))}{len(rows) + 1}"


def _sheet_summary(ws, doc: Doc) -> None:
    """Сводка в лист: заголовки, абзацы, таблицы — те же блоки, что
    в summary.md, чтобы цифры в двух местах не могли разойтись."""
    ws.column_dimensions["A"].width = 44
    for col in "BCDE":
        ws.column_dimensions[col].width = 20
    row = 1
    prev = ""
    for b in doc.blocks:
        if row > 1 and b.kind in ("h1", "h2", "h3", "table", "bars") or (
            b.kind == "p" and prev not in ("p", "")
        ):
            row += 1  # воздух между разделами
        if b.kind in ("h1", "h2", "h3"):
            c = ws.cell(row=row, column=1, value=plain(b.text))
            c.font = {"h1": H1, "h2": H2, "h3": H3}[b.kind]
            row += 1
        elif b.kind in ("p", "li", "quote"):
            text = plain(b.text)
            if b.kind == "li":
                text = f"• {text}"
            c = ws.cell(row=row, column=1, value=text)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if b.kind == "quote":
                c.font = QUOTE
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            # Высота строки под длинный текст: Excel сам её для
            # объединённых ячеек не считает.
            lines = max(1, len(text) // 110 + text.count("\n") + 1)
            ws.row_dimensions[row].height = 15 * lines
            row += 1
        elif b.kind == "table":
            for j, h in enumerate(b.headers, 1):
                c = ws.cell(row=row, column=j, value=h)
                c.font = HEAD_FONT
                c.fill = HEAD_FILL
            row += 1
            for r in b.rows:
                for j, v in enumerate(r, 1):
                    ws.cell(row=row, column=j, value=plain(v) if isinstance(v, str) else v)
                row += 1
        elif b.kind == "bars":
            for j, h in enumerate(("Месяц", "Отзывов"), 1):
                c = ws.cell(row=row, column=j, value=h)
                c.font = HEAD_FONT
                c.fill = HEAD_FILL
            row += 1
            for label, n in b.rows:
                ws.cell(row=row, column=1, value=label)
                ws.cell(row=row, column=2, value=int(n))
                row += 1
        prev = b.kind


def write_workbook(path: Path, doc: Doc, reviews: list[dict], products: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    _sheet_summary(ws, doc)

    _sheet_table(wb.create_sheet("Отзывы"), reviews, FIELDS)

    neg = [
        r for r in reviews
        if str(r.get("rating") or "").isdigit() and int(r["rating"]) <= 3
    ]
    # Лист с негативом нужен всегда, даже пустой: коллега, привыкший
    # к книге, ищет его на третьей вкладке. Пустой лист говорит
    # «негатива в срезе нет» лучше, чем отсутствующий.
    _sheet_table(wb.create_sheet("Негатив"), neg, FIELDS)

    _sheet_table(wb.create_sheet("Карточки"), products, PRODUCT_FIELDS)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
